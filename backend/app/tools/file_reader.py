"""
File Reader Tool
Extracts and summarizes content from TXT, Markdown, PDF, and DOCX files.
"""
import os
from typing import Dict, Any
from app.utils.logger import logger

def read_file_content(file_path: str, max_chars: int = 15000) -> str:
    """
    Reads text content from txt, md, pdf, or docx files.
    Limits text output to max_chars to prevent token overflow.
    """
    if not os.path.exists(file_path):
        return f"Error: File not found at path '{file_path}'"

    ext = os.path.splitext(file_path)[1].lower()
    content = ""
    metadata: Dict[str, Any] = {
        "file_name": os.path.basename(file_path),
        "file_size_bytes": os.path.getsize(file_path),
        "extension": ext
    }

    try:
        if ext in [".txt", ".md", ".py", ".js", ".json", ".csv", ".yaml", ".yml", ".ini", ".cfg", ".html", ".css"]:
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                content = f.read(max_chars)
            metadata["pages"] = 1
            metadata["status"] = "read_success"

        elif ext == ".pdf":
            try:
                import pypdf
                reader = pypdf.PdfReader(file_path)
                num_pages = len(reader.pages)
                metadata["pages"] = num_pages
                
                text_parts = []
                for i in range(num_pages):
                    page_text = reader.pages[i].extract_text() or ""
                    text_parts.append(page_text)
                    current_length = sum(len(p) for p in text_parts)
                    if current_length >= max_chars:
                        break

                content = "\n--- Page Break ---\n".join(text_parts)[:max_chars]
                metadata["status"] = "read_success"
            except ImportError:
                return "Error: pypdf is not installed. Unable to read PDF files."
            except Exception as e:
                return f"Error reading PDF file: {str(e)}"

        elif ext == ".docx":
            try:
                import docx
                doc = docx.Document(file_path)
                text_parts = []
                for para in doc.paragraphs:
                    text_parts.append(para.text)
                    if sum(len(p) for p in text_parts) >= max_chars:
                        break
                
                if sum(len(p) for p in text_parts) < max_chars:
                    for table in doc.tables:
                        for row in table.rows:
                            for cell in row.cells:
                                text_parts.append(cell.text)
                
                content = "\n".join(text_parts)[:max_chars]
                metadata["pages"] = "N/A"
                metadata["status"] = "read_success"
            except ImportError:
                return "Error: python-docx is not installed. Unable to read Word files."
            except Exception as e:
                return f"Error reading Word document: {str(e)}"

        else:
            return f"Error: Unsupported file format '{ext}'. Supported formats: text code, PDF, DOCX."

    except Exception as e:
        logger.error("file_reader_tool_failed", path=file_path, error=str(e))
        return f"Error: Failed to read file content. Details: {str(e)}"

    summary = f"### Document: {metadata['file_name']} ({metadata['extension'].upper()})\n"
    summary += f"**Size:** {metadata['file_size_bytes']} bytes | **Pages:** {metadata['pages']}\n\n"
    summary += "--- START CONTENT ---\n"
    summary += content
    if len(content) >= max_chars:
        summary += "\n\n[TRUNCATED DUE TO SIZE LIMITS]"
    summary += "\n--- END CONTENT ---"
    
    return summary
